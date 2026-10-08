from io import BytesIO
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

POSSIVEIS_SEI = ["NÚMERO SEI", "NUMERO SEI"]
COLUNA_AVISO = "HERDADO DA RELAÇÃO ANTERIOR"
VERDE = PatternFill(fill_type="solid", start_color="E2F0D9", end_color="E2F0D9")
FONTE_VERDE = Font(color="375623")

def norm(v):
    return "" if v is None else str(v).strip()

def headers(ws):
    mapa, orig = {}, {}
    for c in ws[1]:
        if c.value is None:
            continue
        o = norm(c.value)
        k = o.upper()
        if k:
            mapa[k] = c.column
            orig[k] = o
    return mapa, orig

def find_sei(cab):
    for nome in POSSIVEIS_SEI:
        k = nome.upper()
        if k in cab:
            return cab[k]
    raise ValueError("Não foi encontrada a coluna NÚMERO SEI / NUMERO SEI.")

def first_sheet(wb):
    return wb[wb.sheetnames[0]]

def list_columns(content):
    wb = load_workbook(BytesIO(content), data_only=False)
    ws = first_sheet(wb)
    cab, orig = headers(ws)
    find_sei(cab)
    cols = []
    for k, idx in sorted(cab.items(), key=lambda x: x[1]):
        if k not in [x.upper() for x in POSSIVEIS_SEI]:
            cols.append(orig[k])
    return {"aba": ws.title, "colunas": cols}

def copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)

def add_column(ws, title):
    idx = ws.max_column + 1
    cell = ws.cell(1, idx, title)
    if idx > 1:
        copy_style(ws.cell(1, idx-1), cell)
    ws.column_dimensions[cell.column_letter].width = max(18, min(len(title)+4, 35))
    return idx

def process(source_bytes, child_bytes, selected_columns, overwrite=True):
    wb_src = load_workbook(BytesIO(source_bytes), data_only=False)
    wb_child = load_workbook(BytesIO(child_bytes), data_only=False)
    ws_src = first_sheet(wb_src)
    ws_child = first_sheet(wb_child)

    cab_src, orig_src = headers(ws_src)
    cab_child, orig_child = headers(ws_child)
    col_sei_src = find_sei(cab_src)
    col_sei_child = find_sei(cab_child)

    selected = []
    for title in selected_columns:
        key = norm(title).upper()
        if key and key in cab_src and key not in [x.upper() for x in POSSIVEIS_SEI]:
            selected.append(key)
    if not selected:
        raise ValueError("Selecione pelo menos uma coluna para herdar.")

    created = []
    for key in selected:
        if key not in cab_child:
            title = orig_src[key]
            cab_child[key] = add_column(ws_child, title)
            orig_child[key] = title
            created.append(title)

    if COLUNA_AVISO.upper() not in cab_child:
        cab_child[COLUNA_AVISO.upper()] = add_column(ws_child, COLUNA_AVISO)
    col_aviso = cab_child[COLUNA_AVISO.upper()]

    history = {}
    for r in range(2, ws_src.max_row + 1):
        sei = norm(ws_src.cell(r, col_sei_src).value)
        if not sei:
            continue
        data = {}
        for key in selected:
            c = ws_src.cell(r, cab_src[key])
            if norm(c.value):
                data[key] = {"value": c.value, "number_format": c.number_format}
        if data:
            history[sei] = data

    reappeared = inherited = cells = 0
    for r in range(2, ws_child.max_row + 1):
        sei = norm(ws_child.cell(r, col_sei_child).value)
        if not sei or sei not in history:
            continue
        reappeared += 1
        row_fields = []
        for key in selected:
            info = history[sei].get(key)
            if not info:
                continue
            dst = ws_child.cell(r, cab_child[key])
            if (not overwrite) and norm(dst.value):
                continue
            dst.value = info["value"]
            if info.get("number_format"):
                dst.number_format = info["number_format"]
            dst.fill = VERDE
            dst.font = FONTE_VERDE
            row_fields.append(orig_src[key])
            cells += 1
        if row_fields:
            inherited += 1
            ws_child.cell(r, col_aviso).value = "SIM — " + ", ".join(row_fields)
        else:
            ws_child.cell(r, col_aviso).value = "REAPARECEU — sem valor herdado"

    out = BytesIO()
    wb_child.save(out)
    return out.getvalue(), {
        "reapareceram": reappeared,
        "com_heranca": inherited,
        "celulas_herdadas": cells,
        "colunas_criadas": created,
    }
